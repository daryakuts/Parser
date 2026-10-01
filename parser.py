from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

import openpyxl
import pandas
import requests
import xlrd
from bs4 import BeautifulSoup


UNIT_MARKER = "единица измерения"
TARGET_UNIT = "метрическая тонна"


@dataclass(frozen=True)
class TradingResult:
    """Одна строка торгов — запись для таблицы spimex_trading_results."""

    exchange_product_id: str
    exchange_product_name: str
    oil_id: str
    delivery_basis_id: str
    delivery_basis_name: str
    delivery_type_id: str
    volume: Decimal
    total: Decimal
    count: int
    date: date


class SpimexParser:
    def __init__(self, data_start: int, data_end: int, link: str) -> None:
        self.data_start = data_start
        self.data_end = data_end
        self.link = link

    def get_page(self, number_page: int | None = None) -> BeautifulSoup | None:
        parametr = None
        if number_page is not None:
            parametr = {'page': f'page-{number_page}'}
        response = self.open_link(self.link, parametr)
        if response is None:
            return None
        return BeautifulSoup(response.text, 'html.parser')

    def get_links_with_file(self, soup: BeautifulSoup) -> list[dict]:
        elements = soup.find_all('div', class_='accordeon-inner__item')
        results = []

        for item in elements:
            link_tag = None
            for a in item.find_all('a', href=True):
                href = a['href']
                if '.xls' in href.lower() or '/oil_xls/' in href:
                    link_tag = a
                    break
            if link_tag is None:
                continue

            href = link_tag['href']
            date_tag = item.find('span')
            if date_tag is None:
                continue

            m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', date_tag.text)
            if not m:
                continue
            year = int(m.group(3))
            if not (self.data_start <= year <= self.data_end):
                continue

            url = href if href.startswith('http') else 'https://spimex.com' + href
            results.append({'date': year, 'link': url})

        return results

    def get_all_links(self, start_page: int = 19, max_pages: int = 422, empty_streak_limit: int = 5) -> list[dict]:
        all_links: list[dict] = []
        seen: set[str] = set()
        empty_streak = 0

        for page in range(start_page, max_pages + 1):
            soup = self.get_page(page)
            if soup is None:
                print(f'Страница {page}: не загрузилась')
                empty_streak += 1
                if empty_streak >= empty_streak_limit:
                    print(f'{empty_streak} неудач подряд — стоп')
                    break
                continue

            links = self.get_links_with_file(soup)
            if not links:
                print(f'Страница {page}: XLS нет')
                empty_streak += 1
                if empty_streak >= empty_streak_limit:
                    print(f'{empty_streak} пустых подряд — стоп')
                    break
                continue

            empty_streak = 0

            new_links = [l for l in links if l['link'] not in seen]
            for l in new_links:
                seen.add(l['link'])
            all_links.extend(new_links)

            years = [l['date'] for l in links]
            print(
                f'Страница {page}: +{len(new_links)} '
                f'(всего {len(all_links)}), годы {min(years)}..{max(years)}'
            )

            if max(years) < self.data_start:
                print(f'Все годы на странице < {self.data_start} — стоп')
                break

        return all_links

    def download_all_xlsx(self) -> list[dict]:
        links = self.get_all_links()
        print(f'\n=== Всего найдено ссылок: {len(links)} ===')
        for i, item in enumerate(links):
            print(f'{i:3}  {item["date"]!r:40}  {item["link"]}')

        res_file = []
        for i, item in enumerate(links):
            result = self.get_data_with_file(item['link'])
            if result is not None:
                res_file.append({
                    'date': item['date'],
                    'link': item['link'],
                    'sheets': result,
                })
            print(f'[{i + 1}/{len(links)}] обработано')

        print(f'\nИТОГО: ссылок={len(links)}, успешно={len(res_file)}')
        return res_file



    def open_link(self, link: str, parametr: dict | None = None):
        try:
            response = requests.get(link, params=parametr, timeout=300)
        except requests.exceptions.RequestException as e:
            print(f'Ошибка сети при запросе {e}')
            return None
        if response.status_code != 200:
            print(
                f'Ошибка API (Код {response.status_code}). '
                f'Ответ сервера:\n{response.text[:200]}'
            )
            return None
        return response

    def detect_engine(self, data: bytes) -> str:
        if data.startswith(b'\xd0\xcf\x11\xe0'):
            return 'xlrd'
        if data.startswith(b'PK\x03\x04'):
            return 'openpyxl'
        if data.startswith(b'%PDF'):
            raise ValueError('Это PDF')
        if data.startswith(b'<html') or data.startswith(b'<!DOC'):
            raise ValueError('Сайт отдал HTML')
        raise ValueError(f'Неизвестный формат: {data[:8].hex(" ")}')

    def read_excel_file(self, link: str) -> dict:
        response = self.open_link(link)
        content = response.content
        engine = self.detect_engine(content)
        print(f'Движок: {engine}')

        xl_file = pandas.ExcelFile(BytesIO(content), engine=engine)
        print(f'Листы: {xl_file.sheet_names}')
        print('=' * 60)

        all_sheets = {}
        for sheet_name in xl_file.sheet_names:
            print(f"\nЛист: '{sheet_name}'")

            df_raw = pandas.read_excel(
                BytesIO(content),
                sheet_name=sheet_name,
                header=None,
                engine=engine,
            )
            all_sheets[sheet_name] = df_raw

            for i in range(min(20, len(df_raw))):
                row = df_raw.iloc[i]
                row_text = ' | '.join(
                    str(cell) for cell in row if pandas.notna(cell)
                )
                if row_text.strip():
                    print(f"Строка {i}: {row_text[:200]}")
            print('-' * 60)

        return all_sheets

    def get_data_with_file(self, link: str) -> dict | None:
        response = self.open_link(link)
        if response is None:
            return None

        content = response.content
        try:
            engine = self.detect_engine(content)
            print(f"OK URL: {link}")
        except ValueError as e:
            print(f"Error URL: {link} — {e}")
            return None

        all_sheets = {}

        try:
            if engine == 'openpyxl':
                workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
                for sheet_name in workbook.sheetnames:
                    sheet = workbook[sheet_name]
                    rows_list = [
                        list(row) for row in sheet.iter_rows(values_only=True)
                    ]
                    all_sheets[sheet_name] = rows_list
                    print(f"Лист '{sheet_name}': {len(rows_list)} строк")

            elif engine == 'xlrd':
                xl_file = pandas.ExcelFile(BytesIO(content), engine='xlrd')
                for sheet_name in xl_file.sheet_names:
                    df = pandas.read_excel(
                        BytesIO(content),
                        sheet_name=sheet_name,
                        header=None,
                        engine='xlrd',
                    )
                    rows_list = df.values.tolist()
                    rows_list = [
                        [None if pandas.isna(c) else c for c in row]
                        for row in rows_list
                    ]
                    all_sheets[sheet_name] = rows_list
                    print(f"Лист '{sheet_name}': {len(rows_list)} строк")

            total = sum(len(rows) for rows in all_sheets.values())
            print(f"   Всего листов: {len(all_sheets)}, строк: {total}")
            return all_sheets
        except Exception as e:
            print(f"Ошибка парсинга {link}: {e}")
            return None

    def norm(self, value: object) -> str:
        if value is None:
            return ""
        text = (
            str(value)
            .replace("\xa0", " ")
            .replace("ё", "е")
            .strip()
            .lower()
        )
        text = text.replace("обьем", "объем")
        return " ".join(text.split())

    def to_decimal(self, value: object) -> Decimal | None:
        if value is None or value == "":
            return None
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
        if text in {"", "-"}:
            return None
        try:
            return Decimal(text)
        except InvalidOperation:
            return None

    def to_int(self, value: object) -> int | None:
        if value is None or value == "":
            return None
        if isinstance(value, bool):
            return None
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            if value != value:
                return None
            return int(value)
        text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
        text = re.sub(r"[^0-9.\-]", "", text)
        if text in {"", "-", "."}:
            return None
        try:
            return int(Decimal(text))
        except (ValueError, InvalidOperation):
            return None

    def to_str(self, value: object) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def read_rows(self, path: str) -> list[list[object]]:
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(i) for i in range(sheet.nrows)]

    def find_header(self, rows: list[list[object]]) -> tuple[int, dict[str, int]]:
        column_signs: dict[str, tuple[str, ...]] = {
            "exchange_product_id": ("код", "инструмент"),
            "exchange_product_name": ("наименование", "инструмент"),
            "delivery_basis_name": ("базис",),
            "volume": ("объем", "единиц"),
            "total": ("объем", "руб"),
            "count": ("количество", "договор"),
        }

        for i, row in enumerate(rows):
            cells = [self.norm(c) for c in row]
            hits = sum(
                any(k in c for k in keys)
                for keys in column_signs.values()
                for c in cells
            )
            if hits < 2:
                continue

            col_map: dict[str, int] = {}
            for j, c in enumerate(cells):
                for field, keys in column_signs.items():
                    if field in col_map:
                        continue
                    if all(k in c for k in keys):
                        col_map[field] = j
            return i, col_map

        raise ValueError("шапка не найдена")

    def extract_date(self, file_name: str, rows: list[list[object]]) -> date:
        match = re.search(r"(20\d{2}[01]\d[0-3]\d)", file_name)
        if match is None:
            raise ValueError(
                f"Не удалось определить дату торгов: file_name={file_name!r}"
            )
        try:
            return datetime.strptime(match.group(1), "%Y%m%d").date()
        except ValueError:
            raise ValueError(
                f"Найдено {match.group(1)!r}, но это невалидная дата: "
                f"file_name={file_name!r}"
            )

    def parse_sheet(
        self,
        rows: list[list[object]],
        link: str,
    ) -> list[TradingResult]:
        try:
            header_idx, col_map = self.find_header(rows)
        except ValueError as e:
            print(f"Шапка не найдена: {e}")
            return []

        file_name = link.rsplit('/', 1)[-1].split('?')[0]
        trade_date = self.extract_date(file_name, rows)

        records: list[TradingResult] = []
        in_target_block = False

        for i in range(header_idx + 1, len(rows)):
            row = rows[i] if rows[i] is not None else []

            marker_cells = [self.norm(c) for c in row[:6] if c is not None]
            if any(UNIT_MARKER in c for c in marker_cells):
                in_target_block = any(TARGET_UNIT in c for c in marker_cells)
                continue

            if not in_target_block:
                continue

            def cell(field: str):
                j = col_map.get(field)
                if j is None or j >= len(row):
                    return None
                return row[j]

            code_raw = cell("exchange_product_id")
            code = None if code_raw is None else str(code_raw).strip()

            if not code or len(code) < 8:
                continue
            if not re.fullmatch(r"[A-Za-z0-9-]+", code):
                continue
            if code.lower().startswith("итого"):
                continue

            count = self.to_int(cell("count"))
            if count is None or count <= 0:
                continue

            records.append(TradingResult(
                exchange_product_id=code,
                exchange_product_name=self.to_str(cell("exchange_product_name")),
                oil_id=code[:4],
                delivery_basis_id=code[4:7],
                delivery_basis_name=self.to_str(cell("delivery_basis_name")),
                delivery_type_id=code[-1],
                volume=self.to_decimal(cell("volume")) or Decimal("0"),
                total=self.to_decimal(cell("total")) or Decimal("0"),
                count=count,
                date=trade_date,
            ))

        return records

#    def save_to_db(self, ):

if __name__ == '__main__':
    test = SpimexParser(2023, 2026, 'https://spimex.com/markets/oil_products/trades/results/')
